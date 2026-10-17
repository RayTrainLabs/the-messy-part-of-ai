#!/usr/bin/env python3
"""
Phase 6 (part 1) — Compile the per-row scorecard.

Joins scored_rows.jsonl + gap_rows.jsonl by row_id, validates against the 20-column
schema, and writes:
  - aeo_scorecard.csv   (20 columns)
  - aeo_scorecard.xlsx  (20 columns + 3 helper columns at the front, conditional formatting)

Usage:
    python compile_sheet.py --scored output/scored_rows.jsonl \
        --gaps output/gap_rows.jsonl --raw output/raw_responses.jsonl \
        --brand-profile output/brand_profile.json \
        --out-csv output/aeo_scorecard.csv \
        --out-xlsx output/aeo_scorecard.xlsx
"""

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

VALID_POSITIONS = ["Leader", "Challenger", "Alternative", "Niche", "Not competitive", "Not mentioned"]
VALID_ROLES     = ["Primary", "Secondary", "Fallback", "Avoided", "Absent"]
VALID_PRIORITY  = ["Critical", "High", "Medium", "Low"]
VALID_ASSESSMENT = ["Excellent", "Good", "Moderate", "Weak", "Not visible"]

CSV_COLUMNS = [
    "query_tested", "response_summary", "brand_mentioned", "brands_mentioned_list",
    "citation_frequency_score", "share_of_voice_percentage",
    "positional_authority_score", "sentiment_direction_score", "response_relevance_score",
    "market_position_type", "discovery_role", "positioning_narrative", "strength_score",
    "content_gap", "optimization_action", "optimization_rationale",
    "optimization_priority", "optimization_expected_impact",
    "overall_visibility_assessment", "key_insights",
]

XLSX_HELPER_COLUMNS = ["engine", "query_id", "query_category"]


def load_jsonl(path):
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def join_rows(scored, gaps, raw):
    raw_by_id = {r["row_id"]: r for r in raw}
    scored_by_id = {r["row_id"]: r for r in scored if "_scoring_error" not in r}

    out = []
    skipped = []
    for g in gaps:
        rid = g.get("row_id")
        if "_gap_error" in g:
            skipped.append((rid, "gap_error"))
            continue
        s = scored_by_id.get(rid)
        if not s:
            skipped.append((rid, "no_scored_row"))
            continue
        r = raw_by_id.get(rid, {})

        # Defensive cross-check: gap row was generated independently of scored
        # row, but the rubric requires Total Visibility Gap to lead when the
        # brand isn't mentioned. Catch upstream drift before it ships.
        if s.get("brand_mentioned") is False:
            details = g.get("_gap_details") or []
            first_gap = (details[0].get("gap_type") if details else None)
            if first_gap != "Total Visibility Gap":
                skipped.append((rid, f"brand_mentioned=false but first gap is {first_gap!r}, not Total Visibility Gap"))
                continue

        # Normalise brands_mentioned_list (sort + dedupe) instead of rejecting
        # rows where the LLM emitted them in arrival order. Test-run evidence
        # showed ~20% of rows were dropped over a cosmetic ordering issue.
        bm_list = s.get("brands_mentioned_list") or []
        if bm_list:
            bm_list = sorted({b.strip() for b in bm_list if b and b.strip()},
                              key=lambda x: x.lower())
            s["brands_mentioned_list"] = bm_list

        record = {
            # helper columns (XLSX only)
            "engine":        r.get("engine") or s.get("engine"),
            "query_id":      r.get("query_id") or s.get("query_id"),
            "query_category": r.get("category"),

            "query_tested":  s.get("query_tested") or r.get("query_tested"),
            "response_summary": s.get("response_summary"),
            "brand_mentioned":  s.get("brand_mentioned"),
            "brands_mentioned_list": ", ".join(s.get("brands_mentioned_list") or []),
            "citation_frequency_score":  s.get("citation_frequency_score"),
            "share_of_voice_percentage": s.get("share_of_voice_percentage"),
            "positional_authority_score": s.get("positional_authority_score"),
            "sentiment_direction_score":  s.get("sentiment_direction_score"),
            "response_relevance_score":   s.get("response_relevance_score"),
            "market_position_type":       s.get("market_position_type"),
            "discovery_role":             s.get("discovery_role"),
            "positioning_narrative":      s.get("positioning_narrative"),
            "strength_score":             s.get("strength_score"),
            "content_gap":               g.get("content_gap"),
            "optimization_action":       g.get("optimization_action"),
            "optimization_rationale":    g.get("optimization_rationale"),
            "optimization_priority":     g.get("optimization_priority"),
            "optimization_expected_impact": g.get("optimization_expected_impact"),
            "overall_visibility_assessment": g.get("overall_visibility_assessment"),
            "key_insights": " | ".join(g.get("key_insights") or []),
        }
        problems = validate_record(record)
        if problems:
            skipped.append((rid, ", ".join(problems)))
            continue
        out.append(record)
    return out, skipped


def validate_record(rec):
    problems = []
    if not rec.get("query_tested"):
        problems.append("missing query_tested")
    if not isinstance(rec.get("brand_mentioned"), bool):
        problems.append("brand_mentioned not bool")
    if rec.get("market_position_type") not in VALID_POSITIONS:
        problems.append(f"market_position_type invalid")
    if rec.get("discovery_role") not in VALID_ROLES:
        problems.append(f"discovery_role invalid")
    if rec.get("optimization_priority") not in VALID_PRIORITY:
        problems.append(f"optimization_priority invalid")
    if rec.get("overall_visibility_assessment") not in VALID_ASSESSMENT:
        problems.append(f"overall_visibility_assessment invalid")
    sov = rec.get("share_of_voice_percentage")
    if sov is None or not (0.0 <= float(sov) <= 100.0):
        problems.append("share_of_voice_percentage out of range")
    pa = rec.get("positional_authority_score")
    if pa is None or not (0 <= int(pa) <= 100):
        problems.append("positional_authority_score out of range")
    sd = rec.get("sentiment_direction_score")
    if sd is None or not (-50 <= int(sd) <= 50):
        problems.append("sentiment_direction_score out of range")
    rr = rec.get("response_relevance_score")
    if rr is None or not (0 <= int(rr) <= 100):
        problems.append("response_relevance_score out of range")
    cf = rec.get("citation_frequency_score")
    if cf is None or not (0 <= int(cf) <= 100):
        problems.append("citation_frequency_score out of range")
    ss = rec.get("strength_score")
    if ss is None or not (0 <= int(ss) <= 10):
        problems.append("strength_score out of range")
    ki = rec.get("key_insights") or ""
    if ki.count(" | ") != 2:
        problems.append("key_insights must have exactly 3 items separated by ' | '")
    impact = rec.get("optimization_expected_impact") or ""
    # Three timeframe segments separated by semicolons (output-schema.md row 18).
    if impact.count(";") < 2:
        problems.append("optimization_expected_impact must contain three semicolon-separated timeframes")
    return problems


def write_csv(records, path):
    df = pd.DataFrame(records, columns=CSV_COLUMNS)
    df.to_csv(path, index=False, encoding="utf-8")


def write_xlsx(records, path):
    cols = XLSX_HELPER_COLUMNS + CSV_COLUMNS
    df = pd.DataFrame(records, columns=cols)
    df.to_excel(path, index=False, sheet_name="AEO Scorecard")
    _format_xlsx(path)


def _format_xlsx(path):
    wb = load_workbook(path)
    ws = wb.active

    ws.freeze_panes = "A2"
    header_font = Font(bold=True)
    for c in ws[1]:
        c.font = header_font

    headers = {c.value: c.column_letter for c in ws[1]}
    n_rows = ws.max_row

    def col(name):
        return headers.get(name)

    # Numeric color scales / data bars
    if col("share_of_voice_percentage"):
        rng = f"{col('share_of_voice_percentage')}2:{col('share_of_voice_percentage')}{n_rows}"
        ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0,
                                                       end_type="num", end_value=100, color="4F81BD"))
    if col("positional_authority_score"):
        rng = f"{col('positional_authority_score')}2:{col('positional_authority_score')}{n_rows}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=0, start_color="F8696B",
            mid_type="num",   mid_value=50,   mid_color="FFEB84",
            end_type="num",   end_value=100,  end_color="63BE7B"))
    if col("response_relevance_score"):
        rng = f"{col('response_relevance_score')}2:{col('response_relevance_score')}{n_rows}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=0, start_color="F8696B",
            mid_type="num",   mid_value=50,   mid_color="FFEB84",
            end_type="num",   end_value=100,  end_color="63BE7B"))
    if col("sentiment_direction_score"):
        rng = f"{col('sentiment_direction_score')}2:{col('sentiment_direction_score')}{n_rows}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=-50, start_color="F8696B",
            mid_type="num",   mid_value=0,     mid_color="FFFFFF",
            end_type="num",   end_value=50,    end_color="63BE7B"))
    if col("strength_score"):
        rng = f"{col('strength_score')}2:{col('strength_score')}{n_rows}"
        ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0,
                                                       end_type="num", end_value=10, color="63BE7B"))

    # Categorical color fills
    def fill(color): return PatternFill(start_color=color, end_color=color, fill_type="solid")
    enum_colors = {
        "market_position_type": {
            "Leader": "C6EFCE", "Challenger": "DDEBF7", "Alternative": "FFEB9C",
            "Niche": "FCE4D6", "Not competitive": "F8CBAD", "Not mentioned": "F8696B",
        },
        "optimization_priority": {
            "Critical": "F8696B", "High": "F4B084", "Medium": "FFEB9C", "Low": "D9D9D9",
        },
        "overall_visibility_assessment": {
            "Excellent": "63BE7B", "Good": "C6EFCE", "Moderate": "FFEB9C",
            "Weak": "F4B084", "Not visible": "F8696B",
        },
    }
    for col_name, color_map in enum_colors.items():
        if not col(col_name):
            continue
        col_letter = col(col_name)
        for value, color in color_map.items():
            rule = CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill(color))
            ws.conditional_formatting.add(f"{col_letter}2:{col_letter}{n_rows}", rule)

    # Autosize (rough — openpyxl does not support true autosize)
    for c in ws[1]:
        max_len = len(str(c.value or ""))
        col_letter = c.column_letter
        for cell in ws[col_letter][1:]:
            v = str(cell.value) if cell.value is not None else ""
            if len(v) > max_len:
                max_len = min(len(v), 80)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 80)

    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scored", required=True)
    ap.add_argument("--gaps",   required=True)
    ap.add_argument("--raw",    required=True)
    ap.add_argument("--brand-profile", required=True)
    ap.add_argument("--out-csv",  required=True)
    ap.add_argument("--out-xlsx", required=True)
    args = ap.parse_args()

    scored = load_jsonl(args.scored)
    gaps   = load_jsonl(args.gaps)
    raw    = load_jsonl(args.raw)

    records, skipped = join_rows(scored, gaps, raw)
    print(f"Joined {len(records)} valid records; skipped {len(skipped)}.", file=sys.stderr)
    if skipped[:10]:
        print("First 10 skipped:", file=sys.stderr)
        for rid, why in skipped[:10]:
            print(f"  - {rid}: {why}", file=sys.stderr)

    Path(args.out_csv).parent.mkdir(parents=True, exist_ok=True)
    write_csv(records, args.out_csv)
    write_xlsx(records, args.out_xlsx)
    print(f"Wrote {args.out_csv} and {args.out_xlsx}", file=sys.stderr)


if __name__ == "__main__":
    main()
